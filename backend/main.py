import os
import logging
import re
import tempfile
from copy import copy
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

import httpx
import openpyxl
from dotenv import load_dotenv
from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field

# Load the project-root .env: kenzol-demo/.env, not backend/.env.
PROJECT_ROOT = Path(__file__).resolve().parent.parent
load_dotenv(PROJECT_ROOT / ".env")

API_KEY = os.getenv("SAFECUBE_API", "").strip()
PROVIDER_URL = "https://api.sinay.ai/container-tracking/api/v2/shipment"
SEALINES_URL = "https://api.sinay.ai/container-tracking/api/v2/sealines"

EXCEL_FILE_SETTING = os.getenv(
    "EXCEL_FILE",
    "backend/data/customs_exit_with_est.xlsx"
).strip()

EXCEL_PATH = Path(EXCEL_FILE_SETTING)

if not EXCEL_PATH.is_absolute():
    EXCEL_PATH = PROJECT_ROOT / EXCEL_PATH


logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("kenzol-tracking")

app = FastAPI(
    title="Kenzol - Shipping Tracking API",
    version="1.0.0",
    description="Backend proxy for SafeCube/Sinay container tracking.",
)

# Vite normally proxies /api to this backend. CORS is also enabled for local development.
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
    ],
    allow_credentials=False,
    allow_methods=["GET", "POST"],
    allow_headers=["*"],
)

# Carrier names are mapped to SCAC/sealine codes documented by the provider.
# MCL is intentionally not assigned a guessed code. Verify it against GET /sealines
# and set MCL_SCAC in .env before using that carrier explicitly.
CARRIER_SCAC = {
    "MSC": "MSCU",
    "Hapag-Lloyd": "HLCU",
    "CMA CGM": "CMDU",
    "Maersk": "MAEU",
}
if os.getenv("MCL_SCAC", "").strip():
    CARRIER_SCAC["MCL"] = os.getenv("MCL_SCAC", "").strip().upper()

CARRIER_ALIASES = {
    "msc": "MSC",
    "hapag-lloyd": "Hapag-Lloyd",
    "hapag lloyd": "Hapag-Lloyd",
    "hapaglloyd": "Hapag-Lloyd",
    "cma cgm": "CMA CGM",
    "cmacgm": "CMA CGM",
    "maersk": "Maersk",
    "mcl": "MCL",
}


def clean_text(value: Any, fallback: str = "") -> str:
    if value is None:
        return fallback
    if isinstance(value, str):
        return value.strip() or fallback
    return str(value)


def location_name(value: Any) -> str:
    """Turn a provider location object into a readable location string."""
    if isinstance(value, str):
        return value.strip()
    if not isinstance(value, dict):
        return ""
    parts = [
        clean_text(value.get("name")),
        clean_text(value.get("state")),
        clean_text(value.get("country")),
    ]
    return ", ".join(part for part in parts if part)


def route_location(route: Any, key: str) -> str:
    if not isinstance(route, dict):
        return ""
    item = route.get(key)
    if isinstance(item, dict):
        # Route points nest the place under "location"; tolerate flat shapes too.
        return location_name(item.get("location")) or location_name(item)
    return clean_text(item)


def parse_event(event: dict[str, Any], container_number: str = "") -> dict[str, Any]:
    location = event.get("location")
    facility = event.get("facility")
    loc = location_name(location) or location_name(facility)
    vessel = event.get("vessel")
    if isinstance(vessel, dict):
        transport = clean_text(vessel.get("name"))
    else:
        transport = clean_text(event.get("transportType") or event.get("routeType"))

    event_date = clean_text(event.get("date") or event.get("eventDateTime") or event.get("eventDate"))
    # Keep the provider's date string as-is; the frontend formats ISO date/time strings.
    return {
        "id": "|".join([
            event_date,
            clean_text(event.get("eventCode") or event.get("status")),
            clean_text(event.get("description")),
            container_number,
        ]),
        "title": clean_text(event.get("description") or event.get("eventCode") or event.get("status"), "Shipment event"),
        "location": loc,
        "date": event_date,
        "time": "",
        "transport": transport,
        "classifier": clean_text(event.get("eventCode") or event.get("status")),
        "isActual": event.get("isActual"),
        "containerNumber": container_number,
        "eventType": clean_text(event.get("eventType")),
    }


def normalize_provider_response(data: dict[str, Any], requested_bl: str) -> dict[str, Any]:
    metadata = data.get("metadata") or {}
    route = data.get("route") or {}

    # The provider lists every place on the journey in order: the first is the origin and
    # the last is the destination. Route points are the fallback when the list is missing.
    journey = [name for name in (location_name(item) for item in data.get("locations") or []) if name]
    origin = (
        (journey[0] if journey else "")
        or route_location(route, "prepol")
        or route_location(route, "prePol")
        or route_location(route, "pol")
    )
    destination = (
        (journey[-1] if journey else "")
        or route_location(route, "postpod")
        or route_location(route, "postPod")
        or route_location(route, "pod")
    )

    raw_containers = data.get("containers") or []
    containers: list[dict[str, Any]] = []
    all_events: list[dict[str, Any]] = []

    for container in raw_containers:
        if not isinstance(container, dict):
            continue
        number = clean_text(container.get("number") or container.get("containerNumber"))
        containers.append({
            "containerNumber": number,
            "isoCode": clean_text(container.get("isoCode")),
            "sizeType": clean_text(container.get("sizeType")),
            "status": clean_text(container.get("status")),
        })
        for event in container.get("events") or []:
            if isinstance(event, dict):
                all_events.append(parse_event(event, number))

    # Some provider responses may also include shipment-level events.
    for event in data.get("events") or []:
        if isinstance(event, dict):
            all_events.append(parse_event(event))

    # Deduplicate repeated event records while retaining separate container-specific events.
    seen = set()
    events = []
    for event in all_events:
        fingerprint = (
            event.get("id"),
            event.get("title"),
            event.get("location"),
            event.get("date"),
            event.get("containerNumber"),
        )
        if fingerprint not in seen:
            seen.add(fingerprint)
            events.append(event)

    events.sort(key=lambda e: clean_text(e.get("date")))

    status_code = clean_text(metadata.get("shippingStatus"), "UNKNOWN")
    status_labels = {
        "PLANNED": "Planned",
        "IN_TRANSIT": "In transit",
        "DELIVERED": "Delivered",
        "UNKNOWN": "Status unavailable",
    }
    status = status_labels.get(status_code.upper(), status_code.replace("_", " ").title())

    # Estimate ETA only from provider-supplied route dates; never calculate or invent one.
    eta = ""
    pod = route.get("pod") if isinstance(route, dict) else None
    if isinstance(pod, dict):
        eta = clean_text(
            pod.get("eta")
            or pod.get("estimatedArrival")
            or pod.get("estimatedDate")
            or pod.get("date")
        )

    # Return a frontend-friendly response. No provider fields are fabricated.
    return {
        "shipmentNumber": clean_text(metadata.get("shipmentNumber"), requested_bl),
        "carrier": clean_text(metadata.get("sealineName") or metadata.get("sealine"), "Carrier not identified"),
        "status": status,
        "statusCode": status_code,
        "origin": origin or "Not provided by carrier",
        "destination": destination or "Not provided by carrier",
        "eta": eta,
        "lastUpdated": clean_text(metadata.get("updatedAt")),
        "containers": containers,
        "events": events,
        "warnings": metadata.get("warnings") or [],
        "shipmentType": clean_text(metadata.get("shipmentType"), "BL"),
        "sealine": clean_text(metadata.get("sealine")),
        "locations": data.get("locations") or [],
        "vessels": data.get("vessels") or [],
        "coordinates": data.get("coordinates"),
        "routeData": data.get("routeData"),
        "providerRoute": route,
    }


@app.get("/api/health")
async def health() -> dict[str, Any]:
    return {
        "status": "ok",
        "apiKeyConfigured": bool(API_KEY),
        "provider": "SafeCube/Sinay Container Tracking API",
    }


@app.get("/api/carriers")
async def carriers() -> dict[str, Any]:
    return {
        "carriers": [
            {"name": name, "sealine": code, "configured": True}
            for name, code in CARRIER_SCAC.items()
        ],
        "autoDetectAvailable": True,
        "note": "MCL appears only after MCL_SCAC is configured with a verified provider sealine code.",
    }


@app.get("/api/track")
async def track_shipment(
    bl_number: str = Query(..., min_length=5, max_length=80, description="Bill of Lading number"),
    carrier: str | None = Query(default=None, description="Optional carrier name"),
) -> dict[str, Any]:
    if not API_KEY:
        raise HTTPException(
            status_code=500,
            detail="SAFECUBE_API is missing. Add it to the project-root .env file and restart the backend.",
        )

    shipment_number = "".join(bl_number.split())
    if not shipment_number:
        raise HTTPException(status_code=422, detail="Enter a Bill of Lading number.")

    params: dict[str, str] = {
        "shipmentNumber": shipment_number,
        "shipmentType": "BL",
        "route": "true",
        "ais": "false",
    }

    selected_carrier = None
    if carrier and carrier.strip():
        selected_carrier = CARRIER_ALIASES.get(carrier.strip().lower())
        if not selected_carrier:
            raise HTTPException(
                status_code=400,
                detail=f"Unsupported carrier selection: {carrier}. Choose a carrier from the dropdown or use Auto-detect.",
            )
        scac = CARRIER_SCAC.get(selected_carrier)
        if not scac:
            raise HTTPException(
                status_code=400,
                detail=f"A verified SafeCube sealine code is not configured for {selected_carrier}. Use Auto-detect or configure its verified SCAC code in .env.",
            )
        params["sealine"] = scac

    headers = {
        "API_KEY": API_KEY,
        "Accept": "application/json",
    }

    try:
        async with httpx.AsyncClient(timeout=httpx.Timeout(60.0, connect=10.0)) as client:
            response = await client.get(PROVIDER_URL, params=params, headers=headers)
    except httpx.TimeoutException:
        raise HTTPException(
            status_code=504,
            detail="The tracking provider took too long to respond. Please try again shortly.",
        )
    except httpx.RequestError as exc:
        logger.warning("Provider connection error: %s", type(exc).__name__)
        raise HTTPException(
            status_code=502,
            detail="Could not connect to the tracking provider. Check your internet connection and try again.",
        )

    if response.status_code in (401, 403):
        raise HTTPException(
            status_code=502,
            detail="The tracking provider rejected the API key. Check SAFECUBE_API in .env and confirm your API plan/key has access.",
        )
    if response.status_code == 429:
        raise HTTPException(
            status_code=429,
            detail="The tracking provider rate limit was reached. Wait a little before trying again.",
        )
    if response.status_code == 404:
        raise HTTPException(
            status_code=404,
            detail="The tracking endpoint or shipment was not found. Verify the BL number and provider API access.",
        )
    if response.status_code >= 500:
        raise HTTPException(
            status_code=502,
            detail=f"The tracking provider returned an error (HTTP {response.status_code}). Please try again later.",
        )
    if not response.is_success:
        # Do not send raw provider response bodies to the browser; they may contain details.
        raise HTTPException(
            status_code=400,
            detail=f"The tracking provider rejected the request (HTTP {response.status_code}). Verify the BL number, shipment type and carrier.",
        )

    try:
        provider_data = response.json()
    except ValueError:
        raise HTTPException(status_code=502, detail="The tracking provider returned an unreadable response.")

    if not isinstance(provider_data, dict):
        raise HTTPException(status_code=502, detail="The tracking provider returned an unexpected response format.")

    metadata = provider_data.get("metadata") or {}
    provider_message = clean_text(provider_data.get("message") or provider_data.get("error"))
    # Some provider responses encode a logical error in a 200 response.
    if provider_data.get("status") == "error" or provider_data.get("errorCode") is not None:
        raise HTTPException(
            status_code=404,
            detail=provider_message or "The provider could not retrieve this shipment. Check the BL number and carrier.",
        )
    if not metadata and not (provider_data.get("containers") or provider_data.get("events")):
        raise HTTPException(
            status_code=404,
            detail=provider_message or "No shipment details were returned. Check the BL number or try Auto-detect for the carrier.",
        )

    normalized = normalize_provider_response(provider_data, shipment_number)
    if selected_carrier:
        normalized["selectedCarrier"] = selected_carrier
    return normalized


@app.get("/api/sealines")
async def get_sealines() -> dict[str, Any]:
    """Fetch the provider's official sealine list so carrier SCACs can be verified."""
    if not API_KEY:
        raise HTTPException(status_code=500, detail="SAFECUBE_API is missing from the project-root .env file.")
    try:
        async with httpx.AsyncClient(timeout=20.0) as client:
            response = await client.get(SEALINES_URL, headers={"API_KEY": API_KEY, "Accept": "application/json"})
    except httpx.RequestError:
        raise HTTPException(status_code=502, detail="Could not connect to the tracking provider.")
    if response.status_code in (401, 403):
        raise HTTPException(status_code=502, detail="The provider rejected the API key.")
    if not response.is_success:
        raise HTTPException(status_code=502, detail=f"The provider returned HTTP {response.status_code}.")
    try:
        return {"sealines": response.json()}
    except ValueError:
        raise HTTPException(status_code=502, detail="The provider returned an unreadable sealine list.")


def normalize_bl(value: Any) -> str:
    """Normalize Excel BL values, including numeric-looking cells."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return "".join(str(value).split()).strip().upper()


def first_present(*values: Any) -> Any:
    for value in values:
        if value is not None and str(value).strip() != "":
            return value
    return None


def route_field(route: Any, *keys: str) -> Any:
    if not isinstance(route, dict):
        return None
    for key in keys:
        value = route.get(key)
        if isinstance(value, dict):
            candidate = first_present(value.get("etd"), value.get("eta"), value.get("estimatedDate"), value.get("date"))
            if candidate is not None:
                return candidate
        elif value is not None and str(value).strip():
            return value
    return None


def save_workbook(workbook: Any) -> None:
    """Save via a temporary file, then atomically replace the workbook."""
    fd, temp_name = tempfile.mkstemp(prefix=EXCEL_PATH.stem + "_", suffix=".xlsx", dir=EXCEL_PATH.parent)
    os.close(fd)
    try:
        workbook.save(temp_name)
        os.replace(temp_name, EXCEL_PATH)
    finally:
        if os.path.exists(temp_name):
            os.remove(temp_name)


@app.post("/api/update-excel")
async def update_excel(
    bl_number: str = Query(..., min_length=5, max_length=80),
    carrier: str | None = Query(default=None),
) -> dict[str, Any]:
    """
    Fetch current provider data and update the matching BL row in the configured workbook.
    Only tracking columns are changed; unrelated customs/consignee fields and workbook
    formatting are preserved. The workbook is saved back to the same file path.
    """
    if not EXCEL_PATH.exists():
        raise HTTPException(
            status_code=500,
            detail=f"Excel workbook not found at '{EXCEL_PATH}'. Set EXCEL_FILE in the project-root .env.",
        )

    # Fetch the latest data through the same backend tracking path as the UI.
    shipment = await track_shipment(bl_number=bl_number, carrier=carrier)
    provider_route = shipment.get("providerRoute") or {}
    containers = shipment.get("containers") or []
    container_numbers = [
        clean_text(item.get("containerNumber"))
        for item in containers
        if isinstance(item, dict) and clean_text(item.get("containerNumber"))
    ]
    shipping_line = shipment.get("carrier")
    eta = first_present(
        shipment.get("eta"),
        route_field(provider_route, "pod", "postpod", "postPod"),
    )
    etd = route_field(provider_route, "pol", "prepol", "prePol")

    # Open and update in-place via a temporary file, then atomically replace the workbook.
    try:
        workbook = openpyxl.load_workbook(EXCEL_PATH)
    except PermissionError:
        raise HTTPException(
            status_code=423,
            detail="The Excel workbook is open or locked. Close it in Excel and try again.",
        )
    except Exception as exc:
        logger.warning("Could not open workbook: %s", type(exc).__name__)
        raise HTTPException(status_code=500, detail="Could not open the Excel workbook. Check that it is a valid .xlsx file.")

    try:
        sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
        header_map = {}
        for col in range(1, sheet.max_column + 1):
            header = sheet.cell(1, col).value
            if isinstance(header, str) and header.strip():
                header_map[header.strip().casefold()] = col

        bl_col = header_map.get("bl number")
        if not bl_col:
            raise HTTPException(status_code=400, detail="Could not find a 'BL NUMBER' column in the Excel workbook.")

        target_bl = normalize_bl(bl_number)
        matching_rows = [
            row for row in range(2, sheet.max_row + 1)
            if normalize_bl(sheet.cell(row, bl_col).value) == target_bl
        ]

        # Write only fields supplied by the tracking provider. Do not invent customs
        # or consignee details that the provider does not know.
        updates: dict[str, Any] = {
            "Shipping Line": shipping_line,
            "Container No": ", ".join(container_numbers) if container_numbers else None,
            "NO. OF CONTAINERS": len(container_numbers) if container_numbers else None,
            "point of loading": shipment.get("origin"),
            "Destination": shipment.get("destination"),
            "VESSEL ETD": etd,
            "VESSEL ETA": eta,
            "BL NUMBER": shipment.get("shipmentNumber") or bl_number,
        }
        unavailable = {"", "—", "Not provided by carrier", "Carrier not identified", "Status unavailable"}
        updates = {
            name: value for name, value in updates.items()
            if value is not None
            and str(value).strip() not in unavailable
            and name.casefold() in header_map
        }
        if not updates:
            raise HTTPException(status_code=502, detail="The provider returned no usable fields to write to Excel.")

        created_new_row = not matching_rows
        if matching_rows:
            # Update every matching row if the workbook contains duplicate BL entries.
            rows_to_update = matching_rows
        else:
            # The BL is new: append a new row instead of rejecting the update.
            last_data_row = 1
            for candidate_row in range(2, sheet.max_row + 1):
                if any(sheet.cell(candidate_row, col).value is not None for col in range(1, sheet.max_column + 1)):
                    last_data_row = candidate_row
            new_row = max(2, last_data_row + 1)
            rows_to_update = [new_row]

            # Copy formatting only from the previous data row; never copy its data.
            if last_data_row >= 2:
                for col in range(1, sheet.max_column + 1):
                    source_cell = sheet.cell(last_data_row, col)
                    target_cell = sheet.cell(new_row, col)
                    if source_cell.has_style:
                        target_cell._style = copy(source_cell._style)
                    if source_cell.number_format:
                        target_cell.number_format = source_cell.number_format
                    if source_cell.alignment:
                        target_cell.alignment = copy(source_cell.alignment)
                    if source_cell.protection:
                        target_cell.protection = copy(source_cell.protection)
                if sheet.row_dimensions[last_data_row].height:
                    sheet.row_dimensions[new_row].height = sheet.row_dimensions[last_data_row].height

        for row_num in rows_to_update:
            for name, value in updates.items():
                col = header_map.get(name.casefold())
                # Do not erase existing data when the provider doesn't return a value.
                if col and value is not None and str(value).strip() != "":
                    cell = sheet.cell(row_num, col)
                    cell.value = value
                    if name in ("VESSEL ETD", "VESSEL ETA") and isinstance(value, str):
                        cell.number_format = "DD/MM/YYYY"

        save_workbook(workbook)

        return {
            "success": True,
            "message": (
                f"Added BL {bl_number} as a new row in Excel."
                if created_new_row
                else f"Updated existing Excel row(s) for BL {bl_number}."
            ),
            "action": "created" if created_new_row else "overwritten",
            "fileName": EXCEL_PATH.name,
            "rowsUpdated": rows_to_update,
            "fieldsUpdated": list(updates.keys()),
            "shipmentNumber": shipment.get("shipmentNumber"),
        }
    except HTTPException:
        raise
    except PermissionError:
        raise HTTPException(
            status_code=423,
            detail="Could not replace the Excel workbook. Close it in Excel and try again.",
        )
    except Exception as exc:
        logger.exception("Excel update failed")
        raise HTTPException(status_code=500, detail="Failed to save the Excel workbook. No success was reported.")
    finally:
        workbook.close()


# Columns surfaced in the Excel register view, in display order.
REGISTER_COLUMNS = [
    "BL NUMBER",
    "Shipping Line",
    "Container No",
    "NO. OF CONTAINERS",
    "point of loading",
    "Destination",
    "VESSEL ETD",
    "VESSEL ETA",
    "CONSIGNEE NAME",
    "Date",
]
XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Placeholder values used in the workbook where no BL exists yet.
BL_PLACEHOLDERS = {"", "N/A", "NA", "-", "--", "NIL", "NONE", "TBA", "TBC"}


def excel_value(value: Any) -> Any:
    """Make workbook cell values JSON-safe without reformatting provider strings."""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip()
    return value


def require_workbook() -> None:
    if not EXCEL_PATH.exists():
        raise HTTPException(
            status_code=404,
            detail=f"Excel workbook not found at '{EXCEL_PATH.name}'. Set EXCEL_FILE in the project-root .env.",
        )


@app.get("/api/excel")
async def excel_summary() -> dict[str, Any]:
    """Return workbook metadata and the register rows shown in the Excel workspace."""
    require_workbook()
    try:
        workbook = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    except PermissionError:
        raise HTTPException(status_code=423, detail="The Excel workbook is locked. Close it in Excel and try again.")
    except Exception as exc:
        logger.warning("Could not read workbook: %s", type(exc).__name__)
        raise HTTPException(status_code=500, detail="Could not read the Excel workbook. Check that it is a valid .xlsx file.")

    try:
        sheet = workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        headers = next(rows_iter, ()) or ()
        header_map = {
            str(header).strip().casefold(): index
            for index, header in enumerate(headers)
            if isinstance(header, str) and header.strip()
        }
        columns = [name for name in REGISTER_COLUMNS if name.casefold() in header_map]

        rows: list[dict[str, Any]] = []
        for row_number, values in enumerate(rows_iter, start=2):
            if not any(value is not None and str(value).strip() != "" for value in values):
                continue
            record = {
                name: excel_value(values[header_map[name.casefold()]])
                if header_map[name.casefold()] < len(values) else None
                for name in columns
            }
            bl = normalize_bl(record.get("BL NUMBER"))
            rows.append({"row": row_number, "hasBL": bl not in BL_PLACEHOLDERS, "values": record})
    finally:
        workbook.close()

    stat = EXCEL_PATH.stat()
    return {
        "fileName": EXCEL_PATH.name,
        "sheetName": sheet.title,
        "sizeBytes": stat.st_size,
        "modifiedAt": datetime.fromtimestamp(stat.st_mtime, tz=timezone.utc).isoformat(),
        "columns": columns,
        "totalRecords": len(rows),
        "trackedRecords": sum(1 for row in rows if row["hasBL"]),
        "rows": rows,
    }


@app.get("/api/excel/download")
async def download_excel() -> FileResponse:
    """Stream the current workbook so the latest saved updates can be downloaded."""
    require_workbook()
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
    return FileResponse(
        EXCEL_PATH,
        media_type=XLSX_MEDIA_TYPE,
        filename=f"{EXCEL_PATH.stem}_{stamp}.xlsx",
        headers={"Cache-Control": "no-store"},
    )


# Workbook columns shown in the Shipment summary view. CHEQUE AMOUNT is created on first save.
SUMMARY_COLUMNS = ["EST NUMBER", "Exit No", "Agent Name", "VESSEL ETA", "CHEQUE AMOUNT"]
CHEQUE_AMOUNT_HEADER = "CHEQUE AMOUNT"


def bl_tokens(value: Any) -> set[str]:
    """A BL cell may hold several references, e.g. 'DXB1015984C/DXB1015984D'."""
    whole = normalize_bl(value)
    tokens = {normalize_bl(part) for part in re.split(r"[/,;]+", str(value or ""))}
    return {token for token in tokens | {whole} if token not in BL_PLACEHOLDERS}


def load_workbook_or_raise(**kwargs: Any) -> Any:
    require_workbook()
    try:
        return openpyxl.load_workbook(EXCEL_PATH, **kwargs)
    except PermissionError:
        raise HTTPException(status_code=423, detail="The Excel workbook is locked. Close it in Excel and try again.")
    except Exception as exc:
        logger.warning("Could not open workbook: %s", type(exc).__name__)
        raise HTTPException(status_code=500, detail="Could not open the Excel workbook. Check that it is a valid .xlsx file.")


def main_sheet(workbook: Any) -> Any:
    return workbook["Sheet1"] if "Sheet1" in workbook.sheetnames else workbook.active


@app.get("/api/shipment-summary")
async def shipment_summary(
    bl_numbers: str = Query(..., min_length=5, max_length=4000, description="Comma-separated BL numbers"),
) -> dict[str, Any]:
    """Return the customs/agent fields recorded in the workbook for each requested BL."""
    requested = list(dict.fromkeys(
        normalize_bl(part) for part in bl_numbers.split(",") if normalize_bl(part)
    ))
    if not requested:
        raise HTTPException(status_code=422, detail="Enter at least one BL number.")

    workbook = load_workbook_or_raise(read_only=True, data_only=True)
    try:
        sheet = main_sheet(workbook)
        rows_iter = sheet.iter_rows(values_only=True)
        headers = next(rows_iter, ()) or ()
        header_map = {
            str(header).strip().casefold(): index
            for index, header in enumerate(headers)
            if isinstance(header, str) and header.strip()
        }
        bl_index = header_map.get("bl number")
        if bl_index is None:
            raise HTTPException(status_code=400, detail="Could not find a 'BL NUMBER' column in the Excel workbook.")

        # A BL often spans several rows (one per exit). Merge them into one record per BL,
        # keeping each column's distinct values in row order.
        matches = {bl: {"bl": bl, "rows": [], "values": {name: [] for name in SUMMARY_COLUMNS}} for bl in requested}
        for row_number, values in enumerate(rows_iter, start=2):
            if bl_index >= len(values):
                continue
            tokens = bl_tokens(values[bl_index])
            for bl in requested:
                if bl not in tokens:
                    continue
                matches[bl]["rows"].append(row_number)
                for name in SUMMARY_COLUMNS:
                    index = header_map.get(name.casefold())
                    value = excel_value(values[index]) if index is not None and index < len(values) else None
                    merged = matches[bl]["values"][name]
                    if value is not None and value != "" and value not in merged:
                        merged.append(value)
    finally:
        workbook.close()

    return {
        "fileName": EXCEL_PATH.name,
        "columns": SUMMARY_COLUMNS,
        "shipments": [matches[bl] for bl in requested],
    }


def read_header_map(sheet: Any) -> tuple[dict[str, int], int]:
    """Map casefolded headers to 1-based columns; also return the last header column."""
    header_map = {}
    last_header_col = 0
    for col in range(1, sheet.max_column + 1):
        header = sheet.cell(1, col).value
        if isinstance(header, str) and header.strip():
            header_map[header.strip().casefold()] = col
            last_header_col = col
    return header_map, last_header_col


def ensure_column(sheet: Any, header_map: dict[str, int], last_header_col: int, name: str, style_from: str = "") -> int:
    """Return the column for `name`, adding it after the last header (styled like the headers) if missing."""
    col = header_map.get(name.casefold())
    if col:
        return col
    col = last_header_col + 1
    header_cell = sheet.cell(1, col)
    header_cell.value = name
    style_source = sheet.cell(1, header_map.get(style_from.casefold()) or last_header_col)
    if style_source.has_style:
        header_cell._style = copy(style_source._style)
    header_map[name.casefold()] = col
    return col


class ChequeAmountUpdate(BaseModel):
    bl: str = Field(..., min_length=1, max_length=80)
    amount: float | None = None


class ChequeAmountRequest(BaseModel):
    updates: list[ChequeAmountUpdate] = Field(..., min_length=1, max_length=500)


@app.post("/api/shipment-summary/cheque-amounts")
async def save_cheque_amounts(payload: ChequeAmountRequest) -> dict[str, Any]:
    """Write manually entered cheque amounts into the CHEQUE AMOUNT column."""
    workbook = load_workbook_or_raise()
    try:
        sheet = main_sheet(workbook)
        header_map, last_header_col = read_header_map(sheet)

        bl_col = header_map.get("bl number")
        if not bl_col:
            raise HTTPException(status_code=400, detail="Could not find a 'BL NUMBER' column in the Excel workbook.")

        amount_col = ensure_column(sheet, header_map, last_header_col, CHEQUE_AMOUNT_HEADER, style_from="cheque number")

        # The amount belongs to the BL, so it is written to every row that BL appears on,
        # the same way update-excel treats duplicate BL rows.
        amounts_by_bl = {normalize_bl(update.bl): update.amount for update in payload.updates}
        amounts_by_row = {}
        for row_num in range(2, sheet.max_row + 1):
            for bl in bl_tokens(sheet.cell(row_num, bl_col).value) & amounts_by_bl.keys():
                amounts_by_row[row_num] = amounts_by_bl[bl]
        missing = amounts_by_bl.keys() - {
            bl for row_num in amounts_by_row for bl in bl_tokens(sheet.cell(row_num, bl_col).value)
        }
        if missing:
            raise HTTPException(status_code=404, detail=f"BL {', '.join(sorted(missing))} is not in the Excel register.")

        for row_num, amount in amounts_by_row.items():
            cell = sheet.cell(row_num, amount_col)
            cell.value = amount
            if amount is not None:
                cell.number_format = "#,##0.00"

        save_workbook(workbook)
    except HTTPException:
        raise
    except PermissionError:
        raise HTTPException(status_code=423, detail="Could not replace the Excel workbook. Close it in Excel and try again.")
    except Exception:
        logger.exception("Cheque amount update failed")
        raise HTTPException(status_code=500, detail="Failed to save the Excel workbook. No success was reported.")
    finally:
        workbook.close()

    return {
        "success": True,
        "fileName": EXCEL_PATH.name,
        "rowsUpdated": sorted(amounts_by_row),
    }


# Shipments whose vessel ETA is this many days away (or less) need a payment request.
PAYMENT_ALERT_DAYS = 15
PAYMENT_STATUS_HEADER = "PAYMENT STATUS"
PAYMENT_RECEIVED = "Received"


def parse_eta(value: Any) -> date | None:
    """Read workbook ETAs: real dates, ISO strings from the provider, or day-first text like 28/11/26."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value)
    if not text:
        return None
    try:
        return datetime.fromisoformat(text).date()
    except ValueError:
        pass
    match = re.fullmatch(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2}|\d{4})", text)
    if not match:
        return None
    day, month, year = (int(part) for part in match.groups())
    try:
        return date(year + 2000 if year < 100 else year, month, day)
    except ValueError:
        return None


@app.get("/api/payment-alerts")
async def payment_alerts() -> dict[str, Any]:
    """List BLs arriving within PAYMENT_ALERT_DAYS (or already arrived) whose payment is not marked received."""
    workbook = load_workbook_or_raise(read_only=True, data_only=True)
    try:
        rows_iter = main_sheet(workbook).iter_rows(values_only=True)
        headers = next(rows_iter, ()) or ()
        header_map = {
            str(header).strip().casefold(): index
            for index, header in enumerate(headers)
            if isinstance(header, str) and header.strip()
        }
        bl_index = header_map.get("bl number")
        eta_index = header_map.get("vessel eta")
        if bl_index is None or eta_index is None:
            raise HTTPException(status_code=400, detail="The Excel workbook needs 'BL NUMBER' and 'VESSEL ETA' columns.")
        status_index = header_map.get(PAYMENT_STATUS_HEADER.casefold())

        def cell(values: tuple, index: int | None) -> Any:
            return values[index] if index is not None and index < len(values) else None

        today = date.today()
        alerts: dict[str, dict[str, Any]] = {}
        received: set[str] = set()
        for values in rows_iter:
            bl = clean_text(cell(values, bl_index))
            if not bl_tokens(bl):
                continue
            key = normalize_bl(bl)
            if clean_text(cell(values, status_index)).casefold() == PAYMENT_RECEIVED.casefold():
                received.add(key)
                continue
            eta = parse_eta(cell(values, eta_index))
            if eta is None or (eta - today).days > PAYMENT_ALERT_DAYS:
                continue
            # A BL spans several rows; alert once, on its earliest ETA.
            if key not in alerts or eta < date.fromisoformat(alerts[key]["eta"]):
                alerts[key] = {
                    "bl": bl,
                    "eta": eta.isoformat(),
                    "daysLeft": (eta - today).days,
                    "consignee": clean_text(cell(values, header_map.get("consignee name"))),
                    "shippingLine": clean_text(cell(values, header_map.get("shipping line"))),
                }
    finally:
        workbook.close()

    due = sorted((alert for key, alert in alerts.items() if key not in received), key=lambda alert: alert["eta"])
    return {"thresholdDays": PAYMENT_ALERT_DAYS, "alerts": due}


class PaymentReceivedRequest(BaseModel):
    bl: str = Field(..., min_length=1, max_length=200)


@app.post("/api/payment-alerts/received")
async def mark_payment_received(payload: PaymentReceivedRequest) -> dict[str, Any]:
    """Write 'Received' into PAYMENT STATUS on every row of the BL, which stops its alert."""
    workbook = load_workbook_or_raise()
    try:
        sheet = main_sheet(workbook)
        header_map, last_header_col = read_header_map(sheet)
        bl_col = header_map.get("bl number")
        if not bl_col:
            raise HTTPException(status_code=400, detail="Could not find a 'BL NUMBER' column in the Excel workbook.")
        status_col = ensure_column(sheet, header_map, last_header_col, PAYMENT_STATUS_HEADER, style_from="cheque number")

        target = bl_tokens(payload.bl)
        rows = [row for row in range(2, sheet.max_row + 1) if bl_tokens(sheet.cell(row, bl_col).value) & target]
        if not rows:
            raise HTTPException(status_code=404, detail=f"BL {payload.bl} is not in the Excel register.")
        for row in rows:
            sheet.cell(row, status_col).value = PAYMENT_RECEIVED

        save_workbook(workbook)
    except HTTPException:
        raise
    except PermissionError:
        raise HTTPException(status_code=423, detail="Could not replace the Excel workbook. Close it in Excel and try again.")
    except Exception:
        logger.exception("Payment status update failed")
        raise HTTPException(status_code=500, detail="Failed to save the Excel workbook. No success was reported.")
    finally:
        workbook.close()

    return {"success": True, "fileName": EXCEL_PATH.name, "rowsUpdated": rows}
